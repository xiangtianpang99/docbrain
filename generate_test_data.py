import os
from pypdf import PdfWriter
from docx import Document

def create_pdf(filename, text):
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    # pypdf writing text is complex, simplest is just metadata or limited content.
    # Actually, we can use FPDF or ReportLab if installed, but they aren't.
    # Let's stick to simple text file for monitoring test first, 
    # and maybe just try to create a valid DOCX since python-docx is installed and easy.
    # For PDF, I'll just skip complex generation and trust the library if I can't easily make one.
    # Wait, I can use a simple text file renamed for basic content check? No, PDF structure matters.
    # I'll focus on DOCX generation which is easy with python-docx.
    pass

def create_docx(filename, text):
    doc = Document()
    doc.add_paragraph(text)
    doc.save(filename)
    print(f"Created {filename}")

def create_excel(filename):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "SalesData"
    # Header
    ws.append(["Product", "Category", "Revenue", "Date"])
    # Data
    ws.append(["iPhone 15", "Electronics", 8000, "2024-01-01"])
    ws.append(["MacBook Air", "Electronics", 12000, "2024-01-02"])
    ws.append(["Coffee Mug", "Lifestyle", 20, "2024-01-03"])
    ws.append(["Standing Desk", "Furniture", 500, "2024-01-04"])
    
    wb.save(filename)
    print(f"Created {filename}")

if __name__ == "__main__":
    create_docx("data_monitor/test_doc.docx", "The secret password in the DOCX file is: laker.")
    create_excel("data_monitor/sales_test.xlsx")
if __name__ == "__main__":
    create_docx("data_monitor/test_doc.docx", "The secret password in the DOCX file is: ALPHA.")

